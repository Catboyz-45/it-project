"""
สร้างไฟล์สรุปผลการทดสอบ (Excel) จากผลรัน Robot Framework

    ~/.venvs/nestly-robot/bin/python robot-tests/build_report.py

อ่าน robot-tests/results/output.xml แล้วเขียน robot-tests/results/ผลการทดสอบ_TestCase.xlsx
หนึ่งชีตต่อหนึ่งกลุ่มหน้าจอ (ตาม Metadata "Sheet" ของแต่ละไฟล์ .robot) และชีตคำอธิบายท้ายไฟล์

รายละเอียดของแต่ละกรณี (Section / Function / Desc / Type / Data / Steps / Expected / Remark)
อ่านจาก [Documentation] ในไฟล์ .robot ที่เดียว แก้กรณีทดสอบที่ไฟล์ .robot แล้วสร้างรายงานใหม่ ไม่ต้องแก้ที่นี่
"""
from __future__ import annotations

import argparse
import platform
import re
import subprocess
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from robot.api import ExecutionResult
from robot.version import get_version as robot_version

HERE = Path(__file__).resolve().parent
RESULTS = HERE / "results"

PROJECT_NAME = "Nestly ระบบบริหารหอพัก (nestlyspace.tech)"
APP_URL = "http://localhost:3400 (เครื่องทดสอบ · ฐานข้อมูล nestly_robot_test)"
SHEET_ORDER = [
    "Test UI หน้าสาธารณะ",
    "Test UI ผู้เช่า",
    "Test UI เจ้าของหอ",
    "Test UI ผู้ดูแลระบบ",
    "Test สิทธิ์และมือถือ",
]
HEADERS = ["Test case ID", "Section", "Function", "Desc", "P/N", "Data Test", "Test Step",
           "Expected Result", "Status", "Date", "Remark"]
WIDTHS = [13, 18, 20, 40, 6, 30, 36, 40, 9, 19, 44]

FONT = "Tahoma"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
HEADER_FILL = PatternFill("solid", fgColor="C9DAF8")
GROUP_FILL = PatternFill("solid", fgColor="D9D9D9")
COUNT_FILL = PatternFill("solid", fgColor="B6D7A8")

FIELDS = ("Section", "Function", "Desc", "Type", "Data", "Steps", "Expected", "Remark")


def parse_documentation(text: str) -> dict[str, str]:
    """แยก [Documentation] ที่เขียนเป็น "Field: ค่า" ทีละบรรทัด ให้เป็น dict"""
    fields: dict[str, str] = {}
    current = None
    for line in text.splitlines():
        match = re.match(r"^(%s):\s*(.*)$" % "|".join(FIELDS), line.strip())
        if match:
            current = match.group(1)
            fields[current] = match.group(2)
        elif current:
            fields[current] += "\n" + line.strip()
    # ใน .robot เขียนขึ้นบรรทัดใหม่เป็น \n ตัวอักษร แปลงเป็นบรรทัดจริงในเซลล์
    return {key: value.replace("\\n", "\n").strip() for key, value in fields.items()}


def thai_datetime(value: datetime | None) -> str:
    if not value:
        return "-"
    return f"{value.day:02d}/{value.month:02d}/{value.year + 543} {value:%H:%M:%S}"


def chrome_version() -> str:
    try:
        out = subprocess.run(["/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", "--version"],
                             capture_output=True, text=True, timeout=10).stdout
        return out.replace("Google Chrome", "").strip() or "-"
    except (OSError, subprocess.SubprocessError):
        return "-"


def selenium_library_version() -> str:
    try:
        from SeleniumLibrary import __version__
        return __version__
    except ImportError:
        return "-"


def collect(suite, rows: dict[str, list]):
    """เดินทุกชุดทดสอบ เก็บกลุ่ม (หนึ่งไฟล์ .robot) พร้อมกรณีทดสอบ แยกตามชีต"""
    if suite.tests:
        sheet = suite.metadata.get("Sheet", "อื่น ๆ")
        title = suite.doc.splitlines()[0].strip() if suite.doc else suite.name
        pre = suite.metadata.get("Pre-condition", "-").replace("\\n", "\n")
        cases = []
        for test in suite.tests:
            doc = parse_documentation(test.doc)
            test_id = test.name.split(" ", 1)[0]
            remark = doc.get("Remark", "")
            if test.status == "PASS":
                actual = f"ผลจริง: {test.message}" if test.message else ""
            else:
                # ข้อความจาก Robot บอกว่าพลาดที่ขั้นไหน ตัดให้สั้นพอสำหรับเซลล์
                actual = "ไม่ผ่าน: " + test.message.splitlines()[0][:300]
            cases.append({
                "id": test_id,
                "section": doc.get("Section", "-"),
                "function": doc.get("Function", "-"),
                "desc": doc.get("Desc", test.name.split(" ", 1)[-1]),
                "type": doc.get("Type", "P"),
                "data": doc.get("Data", "-"),
                "steps": doc.get("Steps", "-"),
                "expected": doc.get("Expected", "-"),
                "status": "Pass" if test.status == "PASS" else ("Fail" if test.status == "FAIL" else "Not Run"),
                "date": thai_datetime(test.end_time),
                "remark": "\n".join(part for part in (remark, actual) if part),
            })
        rows.setdefault(sheet, []).append({"title": title, "pre": pre, "cases": cases})
    for child in suite.suites:
        collect(child, rows)


def style_cell(cell, *, bold=False, color=None, fill=None, align=TOP_WRAP, border=True):
    cell.font = Font(name=FONT, size=10, bold=bold, color=color)
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if border:
        cell.border = BORDER


def write_sheet(wb: Workbook, name: str, groups: list, created_by: str, environment: str):
    ws = wb.create_sheet(name)
    for index, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    total_rows = 6 + sum(2 + len(group["cases"]) for group in groups)
    info = [("Project name", PROJECT_NAME), ("Created by", created_by), ("Environment", environment), ("Url", APP_URL)]
    for row, (label, value) in enumerate(info, start=1):
        style_cell(ws.cell(row=row, column=1, value=label), bold=True, align=None, border=False)
        style_cell(ws.cell(row=row, column=2, value=value), align=None, border=False)
    counters = [
        (1, "Pass", f'=COUNTIF(I7:I{total_rows},"Pass")'),
        (2, "Fail", f'=COUNTIF(I7:I{total_rows},"Fail")'),
        (3, "Execute", "=I1+I2"),
        (5, "Total", f'=COUNTIF(A7:A{total_rows},"TC-*")'),
    ]
    right = Alignment(horizontal="right")
    for row, label, formula in counters:
        style_cell(ws.cell(row=row, column=8, value=label), fill=COUNT_FILL, align=right)
        style_cell(ws.cell(row=row, column=9, value=formula), fill=COUNT_FILL, align=right)
    for column, header in enumerate(HEADERS, start=1):
        style_cell(ws.cell(row=6, column=column, value=header), bold=True, fill=HEADER_FILL, align=Alignment(vertical="center", wrap_text=True))

    row = 7
    for group in groups:
        for column in range(1, len(HEADERS) + 1):
            style_cell(ws.cell(row=row, column=column, value=group["title"] if column == 1 else None), bold=True, fill=GROUP_FILL)
        row += 1
        for column in range(1, len(HEADERS) + 1):
            value = "Pre-condition" if column == 2 else (group["pre"] if column == 7 else None)
            style_cell(ws.cell(row=row, column=column, value=value), bold=column == 2, color="1155CC")
        row += 1
        for case in group["cases"]:
            values = [case["id"], case["section"], case["function"], case["desc"], case["type"], case["data"],
                      case["steps"], case["expected"], case["status"], case["date"], case["remark"]]
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=column, value=value)
                if column == 9:
                    style_cell(cell, bold=True, color="38761D" if value == "Pass" else "CC0000")
                else:
                    style_cell(cell)
            row += 1
    ws.freeze_panes = "A7"
    return ws


def write_legend(wb: Workbook, summary: list[tuple[str, int, int]]):
    ws = wb.create_sheet("คำอธิบาย")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110
    rows = [
        ("คอลัมน์", "ความหมาย"),
        ("Test case ID", "รหัสกรณีทดสอบ ตรงกับชื่อกรณีทดสอบในไฟล์ .robot (PUB หน้าสาธารณะ / TEN ผู้เช่า / OWN เจ้าของหอ / ADM ผู้ดูแลระบบ / SEC สิทธิ์ / MOB มือถือ)"),
        ("Section", "หน้าจอหรือส่วนของระบบที่ทดสอบ"),
        ("Function", "ฟังก์ชันที่ทดสอบ"),
        ("Desc", "สิ่งที่ระบบควรทำในกรณีนี้"),
        ("P/N", "P = Positive ทดสอบการใช้งานปกติที่ควรสำเร็จ / N = Negative ทดสอบข้อมูลผิดหรือการใช้งานที่ระบบต้องปฏิเสธ"),
        ("Data Test", "ข้อมูลที่ใช้ทดสอบ (<เลขรอบ> คือเลขเวลาที่รัน ใส่ไว้ให้ข้อมูลไม่ซ้ำกันเมื่อรันหลายรอบ)"),
        ("Test Step", "ขั้นตอนที่ Robot Framework ทำแทนผู้ใช้"),
        ("Expected Result", "ผลที่คาดหวัง"),
        ("Status", "ผลจากการรันจริง Pass = ผ่าน / Fail = ไม่ผ่าน"),
        ("Date", "วันเวลาที่กรณีทดสอบนั้นรันเสร็จ (พ.ศ.)"),
        ("Remark", "หมายเหตุของกรณีทดสอบ และผลที่เกิดขึ้นจริงจากการรัน (กรณีไม่ผ่านจะเป็นสาเหตุที่ Robot Framework รายงาน)"),
        ("Pre-condition", "สิ่งที่ต้องทำก่อนเริ่มทุกกรณีในกลุ่มนั้น"),
        (None, None),
        ("ชุดทดสอบ", "ผลรวม"),
        *[(name, f"ผ่าน {passed} / ไม่ผ่าน {failed} / ทั้งหมด {passed + failed} กรณี") for name, passed, failed in summary],
        (None, None),
        ("หลักฐาน", "robot-tests/results/report.html, log.html, output.xml สร้างโดย Robot Framework พร้อมภาพหน้าจอทุกกรณีในโฟลเดอร์ results/screenshots (ชื่อไฟล์ตามรหัสกรณี)"),
        ("วิธีรันซ้ำ", "robot-tests/README.md: เตรียมฐานข้อมูลด้วย prepare.sh เปิดเว็บด้วย start-app.sh แล้วรัน run.sh"),
    ]
    for index, (label, value) in enumerate(rows, start=1):
        header = index == 1 or label == "ชุดทดสอบ"
        if label is None:
            continue
        style_cell(ws.cell(row=index, column=1, value=label), bold=True, fill=HEADER_FILL if header else None)
        style_cell(ws.cell(row=index, column=2, value=value), bold=header, fill=HEADER_FILL if header else None)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", default=str(RESULTS / "output.xml"))
    parser.add_argument("--out", default=str(RESULTS / "ผลการทดสอบ_TestCase.xlsx"))
    parser.add_argument("--created-by", default="นายณัฐภัทร ยืนยง")
    args = parser.parse_args()

    result = ExecutionResult(args.input)
    rows: dict[str, list] = {}
    collect(result.suite, rows)
    environment = (f"Google Chrome {chrome_version()} headless (macOS {platform.mac_ver()[0] or '-'}) · "
                   f"Robot Framework {robot_version()} · SeleniumLibrary {selenium_library_version()}")

    wb = Workbook()
    wb.remove(wb.active)
    summary = []
    for name in SHEET_ORDER + sorted(set(rows) - set(SHEET_ORDER)):
        if name not in rows:
            continue
        write_sheet(wb, name, rows[name], args.created_by, environment)
        cases = [case for group in rows[name] for case in group["cases"]]
        summary.append((name, sum(c["status"] == "Pass" for c in cases), sum(c["status"] != "Pass" for c in cases)))
    write_legend(wb, summary)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    total = sum(p + f for _, p, f in summary)
    passed = sum(p for _, p, _ in summary)
    print(f"เขียน {args.out} แล้ว: ผ่าน {passed} / ทั้งหมด {total} กรณี")


if __name__ == "__main__":
    main()
